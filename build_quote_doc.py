from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"F:\dulizhgan\海外独立站建站服务功能报价表.docx"

NAVY = "17365D"
PALE = "EAF1F8"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=BORDER):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for side in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + side
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    elem = OxmlElement('w:tblHeader')
    elem.set(qn('w:val'), 'true')
    trPr.append(elem)

def set_font(run, size=10.3, bold=False, color=BLACK):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color

def write_cell(cell, text, header=False, center=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, 9.6 if not header else 10, header, RGBColor(255,255,255) if header else BLACK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell); set_cell_border(cell)
    if header: set_cell_shading(cell, NAVY)

def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    widths = [Inches(1.55), Inches(4.25), Inches(1.25)]
    hdr = table.rows[0]
    for i, title in enumerate(('模块','功能内容','建议报价')):
        hdr.cells[i].width = widths[i]; write_cell(hdr.cells[i], title, True, i != 1)
    set_repeat_table_header(hdr)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            write_cell(cells[i], value, False, i != 1)
            if index % 2 == 1: set_cell_shading(cells[i], PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def add_heading(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); set_font(r, 14, True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
for sty in ['Title','Heading 1','Heading 2']:
    styles[sty].font.color.rgb = BLACK

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
r = title.add_run('海外独立站建站服务功能报价表')
set_font(r, 22, True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(18)
r = sub.add_run('适用于海外品牌展示站 电商独立站 与定制增长型项目')
set_font(r, 10.5, False, RGBColor(89,89,89))

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(10)
intro.paragraph_format.line_spacing = 1.35
r = intro.add_run('本报价表用于快速明确独立站项目范围与预算。实际项目按基础建站费、功能模块费及后续维护费组合报价；第三方平台和服务费用由客户承担。')
set_font(r, 10.8)

core = [
('基础展示站','首页、关于我们、产品或服务页、联系页、响应式适配','¥3,000-8,000'),
('企业品牌站','定制视觉、6-15 个页面、表单、基础动画、SEO 结构','¥8,000-20,000'),
('Shopify 商城搭建','主题配置、商品上架、支付及物流配置、基础页面','¥5,000-15,000'),
('WooCommerce 商城搭建','WordPress 与 WooCommerce、支付、商品及订单流程','¥6,000-18,000'),
('定制电商站','定制前后端、后台管理、订单系统','¥25,000-100,000+'),
('UI UX 设计 首页','首页视觉设计','¥1,500-5,000'),
('UI UX 设计 全站','全站页面设计 约 5-10 页','¥5,000-20,000'),
]
add_heading(doc, '一 基础建站与设计')
add_table(doc, core)

commerce = [
('多语言','每增加一种语言，含切换与页面配置，不含翻译','¥800-3,000'),
('多币种','汇率或币种切换、价格展示','¥1,000-4,000'),
('支付接入','Stripe、PayPal 等，每个支付渠道','¥800-3,000'),
('物流配置','运费规则、地区运费、物流追踪','¥1,000-5,000'),
('商品管理','批量导入、规格 SKU、库存、分类筛选','¥1,000-8,000'),
('会员系统','注册登录、个人中心、订单查询','¥2,000-10,000'),
('订阅功能','定期扣款、会员订阅、续费管理','¥3,000-15,000'),
('优惠促销','优惠券、满减、限时折扣、捆绑销售','¥1,000-8,000'),
('评论系统','商品评价、晒单、评价导入','¥800-3,000'),
('预约系统','日历、时段、预约通知、付款预约','¥2,000-10,000'),
('B2B 批发','批发价、分级客户、询盘报价、MOQ','¥4,000-20,000'),
]
add_heading(doc, '二 商城与交易功能')
add_table(doc, commerce)

growth = [
('询盘系统','表单、邮件通知、CRM 表格导出','¥800-4,000'),
('在线客服','Tidio、WhatsApp、Messenger、Live Chat 接入','¥500-2,000'),
('WhatsApp 自动化','浮窗、预填消息、线索标签、自动回复流程','¥1,000-6,000'),
('SEO 基础优化','标题描述、站点地图、Robots、图片 ALT、速度基础优化','¥2,000-8,000'),
('SEO 内容页','博客、案例、落地页模板','¥500-2,000 每页'),
('GA4 与广告像素','GA4、Google Ads、Meta Pixel、转化事件','¥1,000-5,000'),
('邮件营销','Klaviyo 或 Mailchimp 弹窗、订阅、欢迎邮件流','¥1,500-8,000'),
]
add_heading(doc, '三 获客与增长功能')
add_table(doc, growth)

delivery = [
('后台管理','内容、商品、订单、用户等定制后台','¥8,000-50,000+'),
('API 对接','ERP、CRM、供应链、库存或外部服务，每个接口','¥2,000-15,000+'),
('性能优化','CDN、图片优化、缓存、Core Web Vitals','¥2,000-10,000'),
('安全与备份','SSL、备份、权限、基础安全加固','¥1,000-5,000'),
('上线部署','域名、服务器、DNS、Cloudflare、上线检查','¥1,000-5,000'),
('月度维护','小改动、备份、故障处理、插件更新','¥800-5,000 每月'),
]
add_heading(doc, '四 定制开发与交付维护')
add_table(doc, delivery)

add_heading(doc, '五 项目套餐参考')
packages = [
('启动版','外贸企业展示、个人品牌、验证产品','¥5,980-12,800'),
('商业版','正式 Shopify 或 WooCommerce 独立站','¥15,800-39,800'),
('定制增长版','多语言、多币种、营销自动化、定制功能或 API','¥49,800 起'),
]
add_table(doc, packages)

add_heading(doc, '六 报价与合作说明')
notes = [
('第三方费用','域名、服务器、Shopify 月费、付费主题、插件订阅及支付通道手续费，由客户承担。'),
('默认不含项目','翻译、产品拍摄、文案、Logo、广告投放素材，需另行报价。'),
('付款节点','需求确认后收取 50% 定金；测试环境验收后收取 40%；正式上线后收取 10% 尾款。'),
('修改范围','建议首页包含 2-3 轮修改，内页包含 1-2 轮修改；超出范围按 ¥300-1,000 每小时或按页面单独计费。'),
]
for label, text in notes:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.25
    r = p.add_run(label + '：'); set_font(r, 10.5, True)
    r = p.add_run(text); set_font(r, 10.5)

doc.core_properties.title = '海外独立站建站服务功能报价表'
doc.core_properties.subject = '海外独立站建站服务报价'
doc.core_properties.author = ''
doc.save(OUT)
print(OUT)
